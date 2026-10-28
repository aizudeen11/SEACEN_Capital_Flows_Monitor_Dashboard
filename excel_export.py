"""Builds the "Download Figures (Excel)" workbook for the Capital Flow Monitor
dashboard.

Each dashboard chart gets its own sheet: a plain data table (the exact data
the chart is currently showing, after whatever date/region/year filters the
user picked) plus a native openpyxl chart object built from that table. Since
the chart is a real Excel chart object referencing worksheet cells - not a
pasted image - it stays fully editable (colors, chart type, data) once
downloaded.

Chart styling (palette, bottom legend, gridlines, axis lines, no on-chart
titles) is modeled on the hand-built reference workbook "example figures.xlsx"
so exports look consistent with it.

`build_capital_flow_workbook` is the only entry point main.py needs; it takes
the already-filtered DataFrames/selections for each card (same shapes the
card's own `render_plotly` function consumes) and returns an in-memory
workbook ready to stream to the browser.
"""

import io
import itertools

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.legend import Legend
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows

CHART_WIDTH = 24
CHART_HEIGHT = 12
HEADER_FONT = Font(bold=True)

# Palette picked from Excel's own "Standard Colors" swatches, matching the
# manually-colored series in example figures.xlsx, extended enough to cover
# the Policy Rate card's largest region count without repeating.
PALETTE = [
    "00B0F0",  # light blue
    "FFC000",  # gold
    "7030A0",  # purple
    "002060",  # dark blue
    "FF0000",  # red
    "92D050",  # light green
    "0070C0",  # blue
    "00B050",  # green
    "C00000",  # dark red
    "BF8F00",  # dark gold
    "404040",  # dark gray
    "1F4E78",  # dark teal
    "548235",  # dark olive
    "833C00",  # brown
]
NET_LINE_COLOR = "000000"
AXIS_LINE_COLOR = "808080"
GRIDLINE_COLOR = "D9D9D9"
CHART_BG = "FFFFFF"
SERIES_LINE_WIDTH = 28575  # EMU (~2.25pt) - matches example figures.xlsx
AXIS_LINE_WIDTH = 9525  # EMU (~0.75pt)


def _color_cycle():
    return itertools.cycle(PALETTE)


def _write_table(ws, df: pd.DataFrame) -> int:
    """Writes df starting at A1 with a bold header row. Returns the sheet row
    the last data row landed on (header is row 1)."""
    date_cols = [i for i, dt in enumerate(df.dtypes, start=1) if pd.api.types.is_datetime64_any_dtype(dt)]
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for col_idx in date_cols:
        for row_idx in range(2, len(df) + 2):
            ws.cell(row=row_idx, column=col_idx).number_format = "mmm-yyyy"
    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
            12, len(str(col_name)) + 2
        )
    return len(df) + 1


def _style_chart(chart, *, show_legend: bool, value_format: str, pin_axis_to_min: bool):
    """Applies the example workbook's look: no on-chart title (the sheet tab
    carries that), a bottom legend (or none for a single series), muted gray
    axis lines, light horizontal gridlines only, and a white chart area."""
    chart.title = None
    chart.style = None

    if show_legend:
        chart.legend = Legend(legendPos="b")
        chart.legend.overlay = False
    else:
        chart.legend = None

    chart.x_axis.delete = False
    chart.x_axis.title = None
    chart.x_axis.majorGridlines = None
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS_LINE_COLOR, w=AXIS_LINE_WIDTH))
    # By default Excel crosses the category axis at value=0 ("autoZero"), so
    # for any chart whose values straddle zero the category axis - and its
    # labels - end up stranded in the middle of the plot instead of pinned to
    # the edge. Pinning the axis itself to the value axis's minimum (on top of
    # the label-position pin below) fixes that for line charts; for bar
    # charts it instead shifts every bar's baseline off zero, inflating them
    # until they cover the label column, so bars only get the label pin.
    if pin_axis_to_min:
        chart.x_axis.crosses = "min"
    chart.x_axis.tickLblPos = "low"

    chart.y_axis.delete = False
    chart.y_axis.title = None
    chart.y_axis.numFmt = value_format
    chart.y_axis.majorGridlines = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill=GRIDLINE_COLOR, w=AXIS_LINE_WIDTH))
    )
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS_LINE_COLOR, w=AXIS_LINE_WIDTH))

    chart.graphical_properties = GraphicalProperties(solidFill=CHART_BG)
    chart.height, chart.width = CHART_HEIGHT, CHART_WIDTH


def _place_chart(ws, chart, n_cols: int):
    anchor_col_idx = n_cols + 2
    anchor_letter = ws.cell(row=1, column=anchor_col_idx).column_letter
    ws.add_chart(chart, f"{anchor_letter}1")


EVENT_TYPE_LABELS = {"range": "Event period", "line": "Event", "custom": "Custom marker"}


def _add_event_table(ws, start_row: int, events):
    """Writes the active "Show event markers" markers for this card as a
    small table below the data (Date/End Date/Label/Type), so the download
    still carries them even though a native Excel chart can't easily draw the
    dashboard's dashed vline/shaded-band annotations."""
    if not events:
        return
    ws.cell(row=start_row, column=1, value="Event markers").font = Font(bold=True, size=12)
    header_row = start_row + 1
    for col_idx, header in enumerate(["Date", "End Date", "Label", "Type"], start=1):
        ws.cell(row=header_row, column=col_idx, value=header).font = HEADER_FONT
    for i, ev in enumerate(sorted(events, key=lambda e: e["start"]), start=header_row + 1):
        ws.cell(row=i, column=1, value=ev["start"]).number_format = "yyyy-mm-dd"
        if ev["type"] == "range":
            ws.cell(row=i, column=2, value=ev["end"]).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=3, value=ev["label"])
        ws.cell(row=i, column=4, value=EVENT_TYPE_LABELS[ev["type"]])
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 12)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 12)
    ws.column_dimensions["C"].width = max(ws.column_dimensions["C"].width or 0, 45)
    ws.column_dimensions["D"].width = max(ws.column_dimensions["D"].width or 0, 14)


def _line_sheet(wb, sheet_name, wide_df, value_format="#,##0.0", events=None):
    ws = wb.create_sheet(sheet_name)
    last_row = _write_table(ws, wide_df)
    n_cols = len(wide_df.columns)
    n_series = n_cols - 1

    chart = LineChart()
    data = Reference(ws, min_col=2, max_col=n_cols, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    colors = _color_cycle()
    for series in chart.series:
        color = next(colors)
        series.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=color, w=SERIES_LINE_WIDTH)
        )
        series.marker = Marker(symbol="none")
        series.smooth = False

    _style_chart(chart, show_legend=n_series > 1, value_format=value_format, pin_axis_to_min=True)
    _place_chart(ws, chart, n_cols)
    _add_event_table(ws, last_row + 2, events)
    return ws, last_row


def _bar_sheet(
    wb, sheet_name, wide_df, value_format="#,##0.0", bar_type="col", grouping="clustered", events=None
):
    ws = wb.create_sheet(sheet_name)
    last_row = _write_table(ws, wide_df)
    n_cols = len(wide_df.columns)
    n_series = n_cols - 1

    chart = BarChart()
    chart.type = bar_type
    chart.grouping = grouping
    chart.overlap = 100 if grouping == "stacked" else -10
    data = Reference(ws, min_col=2, max_col=n_cols, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    colors = _color_cycle()
    for series in chart.series:
        series.graphicalProperties = GraphicalProperties(solidFill=next(colors))
        # Excel otherwise paints negative bars white ("Invert if negative").
        series.invertIfNegative = False

    _style_chart(chart, show_legend=n_series > 1, value_format=value_format, pin_axis_to_min=False)
    if bar_type == "bar":
        # Plotly's horizontal bars read top-to-bottom in the data's row order
        # (the highest-value region on top); Excel's default is the other
        # way around, so flip the category axis to match.
        chart.x_axis.scaling.orientation = "maxMin"
    _place_chart(ws, chart, n_cols)
    _add_event_table(ws, last_row + 2, events)
    return ws, last_row, n_cols


def _dual_axis_line_sheet(wb, sheet_name, wide_df, rhs_columns, value_format="0.0"):
    """Line chart where the `rhs_columns` series sit on a secondary (right)
    axis, drawn dashed and suffixed " (RHS)" - mirroring the dashboard's
    Headline Inflation chart. The table is written left-axis series first,
    then right-axis ones, so each chart references one contiguous block."""
    first = wide_df.columns[0]
    lhs = [c for c in wide_df.columns[1:] if c not in rhs_columns]
    rhs = [c for c in wide_df.columns[1:] if c in rhs_columns]
    if not lhs:  # only right-axis economies selected: plot them on the main axis
        lhs, rhs = rhs, []
    table = wide_df[[first] + lhs + rhs].rename(columns={c: f"{c} (RHS)" for c in rhs})

    ws = wb.create_sheet(sheet_name)
    last_row = _write_table(ws, table)
    if not lhs:  # nothing selected: keep the (empty) table, skip the chart
        return ws, last_row
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    colors = _color_cycle()

    chart = LineChart()
    chart.add_data(Reference(ws, min_col=2, max_col=1 + len(lhs), min_row=1, max_row=last_row), titles_from_data=True)
    chart.set_categories(cats)
    for series in chart.series:
        series.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=next(colors), w=SERIES_LINE_WIDTH))
        series.marker = Marker(symbol="none")
        series.smooth = False
    _style_chart(chart, show_legend=True, value_format=value_format, pin_axis_to_min=True)

    if rhs:
        right = LineChart()
        right.add_data(
            Reference(ws, min_col=2 + len(lhs), max_col=1 + len(lhs) + len(rhs), min_row=1, max_row=last_row),
            titles_from_data=True,
        )
        right.set_categories(cats)
        for series in right.series:
            series.graphicalProperties = GraphicalProperties(
                ln=LineProperties(solidFill=next(colors), w=SERIES_LINE_WIDTH, prstDash="dash")
            )
            series.marker = Marker(symbol="none")
            series.smooth = False
        right.y_axis.axId = 200
        right.y_axis.crosses = "max"
        right.y_axis.delete = False
        right.y_axis.title = None
        right.y_axis.numFmt = value_format
        right.y_axis.majorGridlines = None
        right.y_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS_LINE_COLOR, w=AXIS_LINE_WIDTH))
        chart += right

    _place_chart(ws, chart, len(table.columns))
    return ws, last_row


def _wide_in_order(df: pd.DataFrame, index_col: str, series_col: str, value_col: str) -> pd.DataFrame:
    """Like `_pivot_long`, but keeps the rows in the order they first appear
    in `df` (e.g. World, Asia-Pacific, China, ... in the forecast charts)
    instead of sorting them."""
    wide = df.pivot_table(index=index_col, columns=series_col, values=value_col, aggfunc="first")
    wide = wide.reindex(df[index_col].unique())
    wide.columns = [str(c) for c in wide.columns]
    return wide.reset_index()


def _pivot_long(df: pd.DataFrame, index_col: str, series_col: str, value_col: str) -> pd.DataFrame:
    """Long-form (index_col, series_col, value_col) -> wide (index_col as the
    first column, one column per distinct series value) - the shape a native
    Excel chart needs, mirroring how the dashboard's Plotly Express `color=`
    charts pick their series."""
    wide = df.pivot_table(index=index_col, columns=series_col, values=value_col, aggfunc="first")
    wide = wide.sort_index()
    wide.columns = [str(c) for c in wide.columns]
    return wide.reset_index()


def _changes_wide(df: pd.DataFrame, selected_years, sort_year) -> pd.DataFrame:
    """Mirrors main.py's `changes_bar_chart`: regions ordered by their value
    in `sort_year` (descending), restricted to the selected years."""
    region_sort = list(df[df["Year"] == sort_year].sort_values(by="Value", ascending=False)["Region"])
    filtered = df[df["Year"].isin(selected_years)]
    wide = filtered.pivot_table(index="Region", columns="Year", values="Value", aggfunc="first")
    wide = wide.reindex(region_sort)
    wide.columns = [str(c) for c in wide.columns]
    return wide.reset_index()


def _tpu_gepu_wide(tpu_df: pd.DataFrame, gepu_df: pd.DataFrame, mode: str):
    if mode == "TPU only":
        return tpu_df.rename(columns={"Value": "TPU"})[["Date", "TPU"]]
    if mode == "GEPU only":
        return gepu_df.rename(columns={"Value": "GEPU"})[["Date", "GEPU"]]
    return pd.merge(
        tpu_df.rename(columns={"Value": "TPU"}),
        gepu_df.rename(columns={"Value": "GEPU"}),
        on="Date",
        how="outer",
    ).sort_values("Date")


POLICY_SHEET_LABELS = {
    "All Region": "All",
    "Selected Region 1": "Reg 1",
    "Selected Region 2": "Reg 2",
}
PU_SHEET_LABELS = {
    "TPU only": "TPU Only",
    "GEPU only": "GEPU Only",
    "Both": "Both",
}


def build_capital_flow_workbook(
    *,
    vix_df: pd.DataFrame,
    tpu_df: pd.DataFrame,
    gepu_df: pd.DataFrame,
    pu_mode: str,
    policy_df: pd.DataFrame | None,
    policy_label: str,
    fx_df: pd.DataFrame,
    fx_years,
    fx_sort_year,
    stock_df: pd.DataFrame,
    stock_years,
    stock_sort_year,
    bond_df: pd.DataFrame,
    bond_years,
    bond_sort_year,
    portfolio_df: pd.DataFrame,
    cds_df: pd.DataFrame | None,
    liquidity_df: pd.DataFrame,
    gdp_df: pd.DataFrame,
    gdp_quarters,
    fsi_df: pd.DataFrame,
    gdp_forecast_df: pd.DataFrame = None,
    inflation_forecast_df: pd.DataFrame = None,
    cpi_df: pd.DataFrame = None,
    cpi_rhs_regions=(),
    vix_events=None,
    pu_events=None,
    policy_events=None,
    portfolio_events=None,
    cds_events=None,
    fsi_events=None,
) -> io.BytesIO:
    """Builds one sheet (data table + native, editable chart) per dashboard
    card, using exactly the filtered data/selections that card is currently
    showing. The `*_events` arguments are the dashboard's active "Show event
    markers" annotations for that card (None/empty when the toggle is off),
    written out as a small table under the chart's data. `gdp_quarters` is in
    chronological order. The policy rate, CDS, forecast and CPI tables are
    optional; their sheets are added only when given (`cpi_rhs_regions` go on
    a right-hand axis)."""
    wb = Workbook()
    wb.remove(wb.active)

    _line_sheet(
        wb, "VIX", vix_df[["Date", "CLOSE"]].rename(columns={"CLOSE": "VIX Close"}), events=vix_events
    )

    pu_label = PU_SHEET_LABELS.get(pu_mode, "Both")
    _line_sheet(wb, f"TPU-GEPU ({pu_label})", _tpu_gepu_wide(tpu_df, gepu_df, pu_mode), events=pu_events)

    if policy_df is not None:
        policy_sheet_label = POLICY_SHEET_LABELS.get(policy_label, "All")
        _line_sheet(
            wb, f"Policy Rate ({policy_sheet_label})", _pivot_long(policy_df, "Date", "Region", "Value"),
            events=policy_events,
        )

    _bar_sheet(
        wb, "FX Changes", _changes_wide(fx_df, fx_years, fx_sort_year),
        value_format="0.00", bar_type="bar",
    )
    _bar_sheet(
        wb, "Stock Index Changes", _changes_wide(stock_df, stock_years, stock_sort_year),
        value_format="0.00", bar_type="bar",
    )
    _bar_sheet(
        wb, "Bond Yield Changes", _changes_wide(bond_df, bond_years, bond_sort_year),
        value_format="0.00", bar_type="bar",
    )

    flows_wide = portfolio_df[["Date", "Portfolio Equity", "Portfolio Debt"]].copy()
    flows_wide["Net Portfolio Flows"] = flows_wide["Portfolio Equity"] + flows_wide["Portfolio Debt"]
    ws_flows, last_row, _ = _bar_sheet(
        wb, "Portfolio Flows", flows_wide.iloc[:, :3], grouping="stacked",
    )
    ws_flows.cell(row=1, column=4, value="Net Portfolio Flows").font = HEADER_FONT
    for i, val in enumerate(flows_wide["Net Portfolio Flows"], start=2):
        ws_flows.cell(row=i, column=4, value=float(val))
    ws_flows.column_dimensions[ws_flows.cell(row=1, column=4).column_letter].width = 20

    net_line = LineChart()
    net_data = Reference(ws_flows, min_col=4, max_col=4, min_row=1, max_row=last_row)
    net_line.add_data(net_data, titles_from_data=True)
    cats = Reference(ws_flows, min_col=1, min_row=2, max_row=last_row)
    net_line.set_categories(cats)
    net_line.series[0].graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=NET_LINE_COLOR, w=SERIES_LINE_WIDTH)
    )
    net_line.series[0].marker = Marker(symbol="none")
    net_line.series[0].smooth = False
    ws_flows._charts[0].legend = Legend(legendPos="b")
    ws_flows._charts[0].legend.overlay = False
    ws_flows._charts[0] += net_line
    _add_event_table(ws_flows, last_row + 2, portfolio_events)

    if cds_df is not None:
        _line_sheet(
            wb, "CDS", _pivot_long(cds_df, "Date", "Region", "Value"), value_format="0.0", events=cds_events
        )

    _bar_sheet(
        wb, "Liquidity", liquidity_df[["TIME_PERIOD", "OBS_VALUE"]].rename(
            columns={"TIME_PERIOD": "Period", "OBS_VALUE": "Liquidity"}
        ),
    )

    gdp_filtered = gdp_df[gdp_df["Quarter"].isin(gdp_quarters)]
    if len(gdp_filtered):
        # Quarters like "1Q2026" sort wrongly as text, so use the given order.
        gdp_wide = _pivot_long(gdp_filtered, "Quarter", "Region", "Value")
        gdp_wide = gdp_wide.set_index("Quarter").reindex(list(gdp_quarters)).reset_index()
        _bar_sheet(wb, "GDP Growth", gdp_wide, value_format="0.00")

    if gdp_forecast_df is not None:
        _bar_sheet(
            wb, "GDP Growth Outlook", _wide_in_order(gdp_forecast_df, "Region", "Year", "Value"), value_format="0.0"
        )
    if inflation_forecast_df is not None:
        _bar_sheet(
            wb, "Inflation Outlook", _wide_in_order(inflation_forecast_df, "Region", "Year", "Value"),
            value_format="0.0",
        )
    if cpi_df is not None:
        _dual_axis_line_sheet(
            wb, "Headline Inflation", _pivot_long(cpi_df, "Date", "Region", "Value"), set(cpi_rhs_regions)
        )

    _line_sheet(wb, "FSI", fsi_df, value_format="0.0", events=fsi_events)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
